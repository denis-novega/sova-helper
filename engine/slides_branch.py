# engine/slides_branch.py
from __future__ import annotations

from pathlib import Path
from typing import List, Dict, Tuple, Optional
import json
import os
import re
import subprocess

# --- внешние либы ---
# PyMuPDF (рендер PDF и извлечение текста)
import fitz  # type: ignore
# OCR
import pytesseract  # type: ignore
from PIL import Image  # type: ignore
# PPTX структура
from pptx import Presentation  # type: ignore

# ---------------------------------------------------------------------
# Конфиг по умолчанию (можно переопределять через ENV)
# ---------------------------------------------------------------------
OCR_LANGS = os.getenv("SOVA_OCR_LANG", "eng+rus")      # какие языки Tesseract
PDF_RENDER_ZOOM = float(os.getenv("SOVA_PDF_ZOOM", "2.0"))  # масштаб рендера (2.0 ≈ 144 DPI)
TESSERACT_CONFIG = os.getenv("SOVA_TESS_CONFIG", "--oem 3 --psm 6")

# ---------------------------------------------------------------------
# Утилиты
# ---------------------------------------------------------------------

def _clean_text(s: str) -> str:
    """Мини-нормализация: убираем хвостовые пробелы, схлопываем множественные пустые строки."""
    s = s.replace("\r", "")
    s = re.sub(r"[ \t]+\n", "\n", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()

def _ensure_dir(p: Path) -> None:
    p.mkdir(parents=True, exist_ok=True)

def _log_append(log_path: Path, msg: str) -> None:
    with log_path.open("a", encoding="utf-8") as f:
        f.write(msg + ("\n" if not msg.endswith("\n") else ""))

def _safe_image_open(path: Path) -> Optional[Image.Image]:
    try:
        return Image.open(path)
    except Exception:
        return None

# ---------------------------------------------------------------------
# LibreOffice: PPT/PPTX -> PDF (для .ppt обязательно; для .pptx только если хотим визуальный PDF-рендер)
# ---------------------------------------------------------------------

def soffice_to_pdf(src: Path, out_dir: Path, err_log: Optional[Path] = None) -> Path:
    """
    Конвертация через LibreOffice в PDF.
    Требует установленный 'soffice' (LibreOffice).
    """
    _ensure_dir(out_dir)
    cmd = ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(src)]
    try:
        if err_log:
            _log_append(err_log, f"--- soffice convert: {' '.join(cmd)}")
            subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.STDOUT)
        else:
            subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except subprocess.CalledProcessError as e:
        raise RuntimeError(f"LibreOffice convert failed ({e.returncode})")
    pdf_path = out_dir / (src.stem + ".pdf")
    if not pdf_path.exists():
        raise RuntimeError("LibreOffice did not produce PDF")
    return pdf_path
# refs: python-pptx docs; LibreOffice headless convert examples. :contentReference[oaicite:0]{index=0}

# ---------------------------------------------------------------------
# PDF: извлечение "сырого" текста и рендер в PNG
# ---------------------------------------------------------------------

def pdf_extract_text_per_page(pdf_path: Path) -> List[str]:
    """Возвращает список текстов по страницам (PyMuPDF get_text('text'))."""
    doc = fitz.open(str(pdf_path))
    out: List[str] = []
    for page in doc:
        txt = page.get_text("text")
        out.append(_clean_text(txt or ""))
    return out
# refs: PyMuPDF docs / discussions. :contentReference[oaicite:1]{index=1}

def render_pdf_to_images(pdf_path: Path, slides_dir: Path, zoom: float = PDF_RENDER_ZOOM) -> List[Path]:
    """
    Рендерит PDF-страницы в PNG (без альфы). Возвращает список путей.
    """
    _ensure_dir(slides_dir)
    doc = fitz.open(str(pdf_path))
    out: List[Path] = []
    for i, page in enumerate(doc, start=1):
        mat = fitz.Matrix(zoom, zoom)
        pix = page.get_pixmap(matrix=mat, alpha=False)
        img_path = slides_dir / f"slide_{i:03d}.png"
        pix.save(str(img_path))
        out.append(img_path)
    return out
# refs: PyMuPDF docs. :contentReference[oaicite:2]{index=2}

# ---------------------------------------------------------------------
# PPTX: извлечение структуры (заголовки, буллеты, заметки)
# ---------------------------------------------------------------------

def pptx_extract_structure(pptx_path: Path) -> List[Dict]:
    prs = Presentation(str(pptx_path))
    slides: List[Dict] = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        bullets: List[Dict] = []
        notes = ""

        # Title
        try:
            if slide.shapes.title and getattr(slide.shapes.title, "text", "").strip():
                title = slide.shapes.title.text.strip()
        except Exception:
            pass

        # Bullets
        for shp in slide.shapes:
            if shp is getattr(slide.shapes, "title", None):
                continue
            if not hasattr(shp, "text_frame") or shp.text_frame is None:
                continue
            tf = shp.text_frame
            for p in tf.paragraphs:
                txt = _clean_text(p.text or "")
                if not txt:
                    continue
                level = getattr(p, "level", 0)
                bullets.append({"text": txt, "level": int(level)})

        # Notes
        try:
            if slide.has_notes_slide and slide.notes_slide and slide.notes_slide.notes_text_frame:
                notes = _clean_text(slide.notes_slide.notes_text_frame.text or "")
        except Exception:
            pass

        slides.append({"index": i, "title": title, "bullets": bullets, "notes": notes})
    return slides
# refs: python-pptx docs: reading text / notes. :contentReference[oaicite:3]{index=3}

# ---------------------------------------------------------------------
# OCR PNG (каждый слайд)
# ---------------------------------------------------------------------

def ocr_image(img_path: Path, langs: str = OCR_LANGS, config: str = TESSERACT_CONFIG) -> str:
    img = _safe_image_open(img_path)
    if img is None:
        return ""
    try:
        txt = pytesseract.image_to_string(img, lang=langs, config=config)  # type: ignore
        return _clean_text(txt or "")
    except Exception:
        return ""

def ocr_slides_pngs(png_paths: List[Path], langs: str = OCR_LANGS, config: str = TESSERACT_CONFIG) -> List[str]:
    out: List[str] = []
    for p in png_paths:
        out.append(ocr_image(p, langs=langs, config=config))
    return out
# refs: pytesseract (madmaze) / Tesseract OCR. :contentReference[oaicite:4]{index=4}

# ---------------------------------------------------------------------
# Сборка outline.md
# ---------------------------------------------------------------------

def build_outline_md(slides: List[Dict]) -> str:
    lines: List[str] = []
    for s in slides:
        idx = s.get("index")
        title = s.get("title") or ""
        bullets = s.get("bullets") or []
        notes = s.get("notes") or ""
        ocr_text = s.get("ocr_text") or ""
        pdf_text = s.get("pdf_text") or ""

        # Заголовок
        if title:
            lines.append(f"## {idx}. {title}")
        else:
            lines.append(f"## Слайд {idx}")

        # Буллеты (с сохранением уровней)
        for b in bullets:
            indent = "  " * min(int(b.get("level", 0)), 4)
            lines.append(f"{indent}- {b.get('text','')}")

        # Заметки докладчика (если есть)
        if notes:
            lines.append(f"\n> Заметки: {notes}\n")

        # Если нет буллетов, но есть pdf_text — добавим его
        if not bullets and pdf_text:
            lines.append(pdf_text)

        # Если нет и pdf_text, и буллетов, но есть OCR — добавим OCR
        if not bullets and not pdf_text and ocr_text:
            lines.append(ocr_text)

        lines.append("")  # пустая строка между слайдами

    return _clean_text("\n".join(lines))

# ---------------------------------------------------------------------
# Главный entrypoint
# ---------------------------------------------------------------------

def process_slides(src_path: Path, workdir: Path, do_ocr: bool = True) -> Dict:
    """
    Универсальный обработчик для .ppt/.pptx/.pdf
    Возвращает dict с ключами: slides, pngs, outline.
    Побочные эффекты: пишет slides/*.png, slides.json, outline.md, document.md
    """
    _ensure_dir(workdir)
    log_path = workdir / "error.log"

    ext = src_path.suffix.lower()
    slides_dir = workdir / "slides"

    # Определяем PDF для рендера
    pdf_path: Optional[Path] = None
    pptx_structure: List[Dict] = []
    pdf_text_pages: List[str] = []

    try:
        if ext in {".ppt", ".pptx"}:
            # Извлекаем структуру из PPTX если возможно
            if ext == ".pptx":
                pptx_structure = pptx_extract_structure(src_path)

            # Для стабильного визуального рендера всё равно получим PDF через LibreOffice
            pdf_path = soffice_to_pdf(src_path, workdir, err_log=log_path)

        elif ext == ".pdf":
            pdf_path = src_path

        else:
            raise ValueError("Unsupported slides format (expected .ppt/.pptx/.pdf)")

        # Рендер PDF в PNG
        pngs = render_pdf_to_images(pdf_path, slides_dir, zoom=PDF_RENDER_ZOOM)

        # Текст из PDF (по страницам)
        try:
            pdf_text_pages = pdf_extract_text_per_page(pdf_path)
        except Exception as e:
            _log_append(log_path, f"[warn] pdf_extract_text_per_page: {e}")
            pdf_text_pages = [""] * len(pngs)

        # OCR всех слайдов
        ocr_pages: List[str] = []
        if do_ocr:
            ocr_pages = ocr_slides_pngs(pngs, langs=OCR_LANGS, config=TESSERACT_CONFIG)
        else:
            ocr_pages = [""] * len(pngs)

        # Склейка структуры по индексам
        count = len(pngs)
        slides: List[Dict] = []
        for i in range(count):
            # если есть pptx_structure — берём соответствующий объект, иначе создаём пустой
            base = pptx_structure[i] if i < len(pptx_structure) else {
                "index": i + 1,
                "title": "",
                "bullets": [],
                "notes": "",
            }

            # добавим извлечённый текст PDF и OCR
            base["pdf_text"] = pdf_text_pages[i] if i < len(pdf_text_pages) else ""
            base["ocr_text"] = ocr_pages[i] if i < len(ocr_pages) else ""
            # относительный путь до PNG внутри workdir (удобно для фронта)
            rel_png = (slides_dir / f"slide_{i+1:03d}.png").relative_to(workdir)
            base["image"] = str(rel_png).replace("\\", "/")

            slides.append(base)

        # outline + сохранение артефактов
        outline_md = build_outline_md(slides)
        (workdir / "slides.json").write_text(
            json.dumps(slides, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        (workdir / "outline.md").write_text(outline_md if outline_md else "(Пусто)", encoding="utf-8")
        # для совместимости
        (workdir / "document.md").write_text(outline_md if outline_md else "(Пусто)", encoding="utf-8")

        return {
            "slides": slides,
            "pngs": [str(p) for p in (slides_dir.glob("*.png"))],
            "outline": outline_md,
        }

    except Exception as e:
        _log_append(log_path, f"[error] {type(e).__name__}: {e}")
        raise
