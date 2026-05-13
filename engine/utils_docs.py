from pathlib import Path
from pdfminer.high_level import extract_text
from pptx import Presentation

def parse_pdf(pdf_path: Path) -> str:
    try:
        return extract_text(str(pdf_path))
    except Exception:
        return ""

def parse_pptx(ppt_path: Path) -> str:
    out = []
    prs = Presentation(str(ppt_path))
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                s = shape.text.strip()
                if s:
                    texts.append(s)
        if texts:
            out.append(f"[Слайд {i}]\n" + "\n".join(texts))
    return "\n\n".join(out)

def build_markdown(lecture_type: str, transcript_text: str, segments, slides_text: str, pdf_text: str) -> str:
    lines = [f"# Документ-лекция ({lecture_type})\n"]
    if lecture_type == "summary":
        lines += [
            "## Краткое резюме",
            "- (сюда ляжет авто-саммари)",
            "## Ключевые термины",
            "- (сюда лягут термины)",
            "## Основной конспект",
            transcript_text[:2000] or "(нет транскрипта)",
        ]
    elif lecture_type == "transcript":
        lines += ["## Транскрипт", transcript_text or "(нет транскрипта)"]
    elif lecture_type == "cheatsheet":
        lines += ["## Шпаргалка", "- (формулы/правила)", "## Q→A", "- Q: ?  A: ?"]
    elif lecture_type == "slides":
        lines += ["## Слайды с тезисами", slides_text or "(нет слайдов)"]

    if pdf_text:
        lines += ["\n## Извлечённый текст PDF", pdf_text[:2000]]

    return "\n".join(lines)
